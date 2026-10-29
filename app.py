"""
app.py - Bank Merger Regulatory Check (Streamlit).

Run:    python3 -m streamlit run app.py
Layout: step list on the left (green arrows show where you are and what's next).
        Home (choose the deal) -> 1 At a glance -> 2 Capital ratios -> 3 Leverage & liabilities
        -> 4 Canadian buyer -> 5 Summary & next steps.  Plus a "Size examples" reference page.
Rule:   engine.py calculates and decides; ai.py reads and explains; this file only shows things.
"""

import datetime
import hashlib
import io
import json
import os
import re

import pandas as pd
import streamlit as st

import ai
import engine as E
from library import Library

try:  # local development: read API keys from a .env file (never upload it)
    from dotenv import load_dotenv
    load_dotenv()
except ImportError:
    pass

st.set_page_config(page_title="Bank Merger Regulatory Check", page_icon="🏦", layout="wide")
HERE = os.path.dirname(os.path.abspath(__file__))
TEST = os.path.join(HERE, "test_data")
SMALL_BANKS = "BHCF20260930.txt"   # the Fed's bulk FR Y-9C file for 18 smaller U.S. holding companies
US_ARM_FILES = {                    # FR Y-9C reports of the Canadian banks' U.S. holding companies
    "TD": ["TD_US_FRY9C_2026Q1.pdf", "TD_US_FRY9C_2026Q2.pdf"], "BMO": ["BMO_US_FRY9C_2026Q2.pdf"],
    "RBC": ["RBC_US_FRY9C_2026Q2.pdf"], "CIBC": ["CIBC_US_FRY9C_2026Q2.pdf"]}
PILLAR3 = {"RBC": ("RBC_Pillar3_Q3_2026.pdf", "RBC Pillar 3 report (Q3 2026)"),
           "BMO": ("BMO_RegCapital_Q3_2026.pdf", "BMO capital disclosure (Q3 2026)"),
           "TD": ("TD_RegDisclosure_Q3_2023.pdf", "TD disclosure (Q3 2023 — older layout)")}
with open(os.path.join(HERE, "library", "us_structure.json")) as _fh:
    STRUCTURE = json.load(_fh)   # built once by build_structure.py from the Fed's NIC ownership data
CANADIAN = list(STRUCTURE)       # TD, RBC, BMO, CIBC

PAGES = ["Home", "1  At a glance", "2  Capital ratios", "3  Leverage & liabilities", "4  Canadian buyer", "5  Summary & next steps"]
EXTRA = "Size examples"


# Streamlit reads "$...$" as maths. Escape dollar signs in plain text (raw HTML blocks are left alone).
def _escape_dollars(fn):
    def wrapper(body="", *args, **kwargs):
        # Block-level HTML (<table>, <div>, <p>) is not run through markdown, so it is left alone;
        # everything else - including inline <span> lines - gets its dollar signs escaped.
        if isinstance(body, str) and not re.match(r"\s*<(table|div|style|p)[\s>]", body):
            body = re.sub(r"(?<!\\)\$", r"\\$", body)
        return fn(body, *args, **kwargs)
    return wrapper


for _n in ("markdown", "info", "success", "warning", "error", "caption", "write"):
    setattr(st, _n, _escape_dollars(getattr(st, _n)))

# --------------------------------------------------------------------------------------
# Look and feel: white, calm, generous spacing, two typefaces
# --------------------------------------------------------------------------------------
st.markdown("""
<style>
@import url('https://fonts.googleapis.com/css2?family=Inter:wght@400;500;600;700&family=Source+Serif+4:opsz,wght@8..60,600;8..60,700&display=swap');
:root { --ink:#1a1d23; --muted:#6b7280; --line:#e7e9ee; --navy:#1f3a93; --green:#1e7d4f; --ai:#6c4ee0; }
.stApp, [data-testid="stAppViewContainer"], [data-testid="stHeader"], [data-testid="stMain"] { background:#ffffff !important; color:var(--ink); }
[data-testid="stSidebar"], [data-testid="stSidebar"] > div { background:#f7f8fa !important; }
.stApp p, .stApp li, .stApp label, .stApp td, .stApp th, .stApp input, .stApp textarea, .stMarkdown {
  font-family:'Inter', system-ui, sans-serif !important; color:var(--ink); }
.stApp button, .stApp button p { font-family:'Inter', system-ui, sans-serif !important; }
.stApp button[kind="primary"], .stApp button[kind="primary"] p, .stApp button[kind="primaryFormSubmit"], .stApp button[kind="primaryFormSubmit"] p { color:#ffffff !important; }
.stApp button[kind="secondary"] p { color:var(--ink); }
.stApp p, .stApp li { font-size:17px; line-height:1.6; }
h1, h2, h3 { font-family:'Source Serif 4', Georgia, serif !important; color:var(--ink) !important; letter-spacing:-0.01em; }
h1 { font-size:2.4rem !important; font-weight:700 !important; }
h2 { font-size:1.75rem !important; font-weight:700 !important; margin-top:0.2rem !important; }
h3 { font-size:1.25rem !important; font-weight:600 !important; }
.block-container { padding-top:2.6rem; max-width:1100px; }
.lede { font-size:19px; color:#3c4250; max-width:760px; line-height:1.55; }
.muted { color:var(--muted) !important; font-size:14.5px; }
.pill { display:inline-block; padding:3px 11px; border-radius:99px; font-weight:600; font-size:13.5px; white-space:nowrap; }
table.nice { width:100%; border-collapse:collapse; margin:4px 0 18px; }
table.nice th { text-align:left; font-weight:600; font-size:13px; text-transform:uppercase; letter-spacing:.04em; color:var(--muted);
  border-bottom:1px solid var(--line); padding:10px 12px; background:#fff; }
table.nice td { padding:12px; border-bottom:1px solid var(--line); vertical-align:top; font-size:15.5px; background:#fff; color:var(--ink); }
table.nice td.src { color:var(--muted); font-size:13.5px; }
.stat { border:1px solid var(--line); border-radius:14px; padding:14px 16px; background:#fff; }
.stat .n { font-size:30px; font-weight:700; font-family:'Inter'; }
.stat .l { font-size:14px; color:var(--muted); }
.card { border:1px solid var(--line); border-radius:16px; padding:18px 20px; background:#fff; }
.step-preview { display:flex; gap:14px; align-items:flex-start; padding:10px 0; border-bottom:1px solid var(--line); }
.step-preview .num { width:30px; height:30px; border-radius:50%; background:#eef2ff; color:var(--navy); font-weight:700;
  display:flex; align-items:center; justify-content:center; flex-shrink:0; }
div.stButton > button { border-radius:10px; font-weight:600; }
/* left-hand step list */
.st-key-nav div.stButton > button > div { justify-content:flex-start; }
.st-key-nav div.stButton > button { justify-content:flex-start; text-align:left; border:1px solid transparent; background:transparent;
  color:var(--ink); padding:.55rem .8rem; font-weight:500; }
.st-key-nav div.stButton > button:hover { background:#eceef3; }
.st-key-nav div.stButton > button[kind="primary"] { background:var(--navy); color:#fff; font-weight:700; }
.arrows { color:var(--green); text-align:center; font-weight:800; line-height:.62; font-size:16px; margin:2px 0 4px; letter-spacing:0; }
.gap { height:6px; border-left:2px solid #dfe3ea; margin-left:26px; }
/* AI answers */
[class*="st-key-ai_"] { border-left:4px solid var(--ai); background:#f8f6ff; border-radius:10px; padding:14px 18px 6px; margin:6px 0 14px; }
.ai-tag { display:inline-block; background:#efeafe; color:#5a3fc0; font-size:12.5px; font-weight:600; padding:2px 10px; border-radius:99px; }
.code-tag { display:inline-block; background:#e7f4ec; color:#1e7d4f; font-size:12.5px; font-weight:600; padding:2px 10px; border-radius:99px; }
/* force a light look for every control, whatever theme the browser or server uses */
.stApp button[kind="secondary"], .stApp button[kind="secondaryFormSubmit"], [data-testid="stPopover"] button {
  background:#ffffff !important; color:var(--ink) !important; border:1px solid #d0d5dd !important; }
.stApp button[kind="secondary"]:hover, [data-testid="stPopover"] button:hover { border-color:var(--navy) !important; color:var(--navy) !important; }
.stApp button[kind="secondary"] p, [data-testid="stPopover"] button p { color:inherit !important; }
.stApp button[kind="primary"], .stApp button[kind="primaryFormSubmit"] { background:var(--navy) !important; border:1px solid var(--navy) !important; }
.stApp button[kind="tertiary"], .stApp button[kind="tertiary"] p { color:var(--navy) !important; background:transparent !important; text-decoration:underline; }
[data-baseweb="select"] > div, [data-baseweb="input"], [data-baseweb="input"] input, [data-baseweb="textarea"], textarea,
[data-testid="stNumberInputContainer"], [data-testid="stNumberInputContainer"] input, [data-testid="stNumberInputContainer"] button {
  background:#f5f7fa !important; color:var(--ink) !important; border-color:#d0d5dd !important; }
[data-baseweb="select"] span, [data-baseweb="select"] div { color:var(--ink) !important; }
[data-baseweb="popover"] ul, [data-baseweb="popover"] li, [data-baseweb="menu"] { background:#ffffff !important; color:var(--ink) !important; }
[data-testid="stFileUploaderDropzone"] { background:#f5f7fa !important; color:var(--ink) !important; }
[data-testid="stFileUploaderDropzone"] * { color:var(--ink) !important; }
/* checkbox ticks and radio dots: navy instead of the default red */
[data-testid="stCheckbox"] div:has(> svg polyline) { background-color:var(--navy) !important; border-color:var(--navy) !important; }
[data-testid="stRadioOption"]:has(input:checked) div:has(> div) { background-color:var(--navy) !important; border-color:var(--navy) !important; }
/* the left-hand step list stays plain text, not boxes */
.st-key-nav div.stButton > button[kind="secondary"] { background:transparent !important; border:1px solid transparent !important; }
.st-key-nav div.stButton > button[kind="secondary"]:hover { background:#eceef3 !important; }
[data-testid="stExpander"] details, [data-testid="stExpander"] summary { background:#ffffff !important; color:var(--ink) !important; }
[data-testid="stVerticalBlockBorderWrapper"] { background:#ffffff; }
.howto { display:flex; gap:12px; align-items:flex-start; margin:10px 0; }
.howto .ic { font-size:20px; line-height:1.4; }
.howto b { display:block; font-size:16px; }
.howto span { color:var(--muted); font-size:15px; }
.diagram { display:flex; flex-direction:column; align-items:center; gap:6px; margin:10px 0 18px; }
.box { border:2px solid var(--line); border-radius:12px; padding:10px 16px; text-align:center; background:#fff; min-width:230px; }
.box.can { border-color:#c8102e; } .box.ihc { border-color:var(--navy); background:#f3f6ff; } .box.new { border-color:var(--green); background:#f1faf5; }
.box small { color:var(--muted); display:block; }
.row { display:flex; gap:12px; flex-wrap:wrap; justify-content:center; }
.down { color:var(--green); font-weight:800; }
</style>
""", unsafe_allow_html=True)


# --------------------------------------------------------------------------------------
# State and small helpers
# --------------------------------------------------------------------------------------
@st.cache_resource
def load_library():
    return Library()


LIB = load_library()
S = st.session_state
S.setdefault("page", "Home")
S.setdefault("saved", {})
S.setdefault("ai_cache", {})
S.setdefault("chat", [])
S.setdefault("chat_i", 0)


def saved(name, default=None):
    return S["saved"].get(name, default)


def keep(name, value):
    S["saved"][name] = value
    return value


def go(page):
    S["page"] = page


def fp(*parts):
    return hashlib.md5(json.dumps(parts, sort_keys=True, default=str).encode()).hexdigest()[:12]


def ready():
    return bool(S.get("banks") and saved("acq") and saved("tgt") and saved("acq") != saved("tgt"))


def nice_date(q):
    try:
        return datetime.datetime.strptime(q, "%Y%m%d").strftime("%d %b %Y")
    except (TypeError, ValueError):
        return str(q)


ACRONYMS = {"Rbc": "RBC", "Bmo": "BMO", "Td": "TD", "Cibc": "CIBC", "Us": "US", "Usa": "USA", "(Usa)": "(USA)", "Llc": "LLC",
            "Lp": "LP", "N.a.": "N.A.", "Wafd,": "WaFd,", "Of": "of", "The": "the", "And": "and"}


def nice_name(name):
    words = (name or "").strip().title().split()
    out = [ACRONYMS.get(w, w) for w in words]
    if out:
        out[0] = out[0][0].upper() + out[0][1:]
    return " ".join(out)


def pill(status):
    icon, label, colour = E.STATUS[status]
    return f'<span class="pill" style="background:{colour}1a;color:{colour}">{icon}&nbsp; {label}</span>'


def pct(v):
    return "n/a" if v is None else f"{v:.2f}%"


def table(rows, cols, src_col=None):
    """Clean HTML table. rows = list of dicts; values may contain HTML (pills)."""
    head = "".join(f"<th>{c}</th>" for c in cols)
    body = "".join("<tr>" + "".join(f'<td class="{"src" if c == src_col else ""}">{r.get(c, "")}</td>' for c in cols) + "</tr>" for r in rows)
    st.markdown(f'<table class="nice"><thead><tr>{head}</tr></thead><tbody>{body}</tbody></table>', unsafe_allow_html=True)


# --------------------------------------------------------------------------------------
# Left-hand side: step list with green arrows, then API credentials
#   >>> FOR SUBMISSION: replace the credentials block with the required openai_credentials module
#       (https://github.com/stoneman-mls/openai_credentials).
# --------------------------------------------------------------------------------------
with st.sidebar:
    st.markdown("### 🏦 Merger check")
    nav = st.container(key="nav")
    current = PAGES.index(S["page"]) if S["page"] in PAGES else None
    with nav:
        for i, p in enumerate(PAGES):
            done = current is not None and i < current
            label = ("✓  " if done and i > 0 else "") + p
            st.button(label, key=f"nav_{i}", on_click=go, args=(p,), type="primary" if S["page"] == p else "secondary",
                      width="stretch")
            if i < len(PAGES) - 1:
                if current == i:
                    st.markdown('<div class="arrows">⌄<br>⌄<br>⌄</div>', unsafe_allow_html=True)
                else:
                    st.markdown('<div class="gap"></div>', unsafe_allow_html=True)
        st.markdown('<div style="height:10px"></div>', unsafe_allow_html=True)
        st.button("📏  " + EXTRA, key="nav_extra", on_click=go, args=(EXTRA,),
                  type="primary" if S["page"] == EXTRA else "secondary", width="stretch")
    st.divider()
    with st.expander("🔑 API key", expanded=not bool(S.get("api_key"))):
        provider = st.radio("Provider", ["Anthropic (Claude)", "OpenAI"], horizontal=True,
                            index=0 if saved("provider", "Anthropic (Claude)").startswith("Anthropic") else 1)
        keep("provider", provider)
        env_key = os.environ.get("ANTHROPIC_API_KEY" if provider.startswith("Anthropic") else "OPENAI_API_KEY", "")
        key = st.text_input("Key", type="password", value=env_key, help="Held only in this browser session; never saved.")
        model = st.text_input("Model", value=ai.DEFAULT_MODELS[provider])
        S["provider"], S["api_key"], S["model"] = provider, key, model
    counter = st.empty()


# --------------------------------------------------------------------------------------
# AI display and caching
# --------------------------------------------------------------------------------------
def run_ai(cache_key, fn, label="Thinking…"):
    if cache_key in S["ai_cache"]:
        return S["ai_cache"][cache_key]
    if not ai.ready():
        st.info("Add your API key in the left-hand panel to see the AI's explanation here. Everything else works without it.")
        return None
    with st.spinner(label):
        try:
            out = fn()
        except Exception as exc:
            st.error(f"The AI call did not work: {exc}")
            return None
    S["ai_cache"][cache_key] = out
    return out


def show_ai(result, title, key):
    if not result:
        return
    with st.container(key=f"ai_{key}"):
        chk = result.get("check")
        note = ""
        if chk:
            note = (f"cites {len(chk['cited'])} regulation passage(s) — all verified" if chk["has_citations"] and not chk["invalid"]
                    else "⚠️ unverified citations were removed" if chk["invalid"] else "no citations")
        st.markdown(f'<span class="ai-tag">✨ AI explanation</span> <span class="muted">&nbsp;{title}{" · " + note if note else ""}</span>',
                    unsafe_allow_html=True)
        st.markdown(result["text"])
        if result.get("passages"):
            with st.expander(f"Sources the AI was allowed to use ({len(result['passages'])})"):
                for p in result["passages"]:
                    st.markdown(f"**[{p['id']}] {p['cite']}** — {p['title']}  \n<span class='muted'>{p['status']} · as of {p['as_of']}</span>",
                                unsafe_allow_html=True)
                    st.caption(p["text"][:700] + ("…" if len(p["text"]) > 700 else ""))


def chat_box(page, facts):
    """One running Q&A. The box clears after each question; ◀ ▶ step through earlier answers."""
    st.markdown("### Ask a question")
    st.markdown('<span class="muted">Answers come only from the regulations loaded into this app, and a second AI call checks each one.</span>',
                unsafe_allow_html=True)
    with st.form(key=f"chatform_{page}", clear_on_submit=True, border=False):
        c1, c2 = st.columns([5, 1])
        q = c1.text_input("Question", label_visibility="collapsed", placeholder="e.g. What does the stress capital buffer do?")
        sent = c2.form_submit_button("Ask", type="primary", width="stretch")
    if sent and q.strip():
        if not ai.ready():
            st.warning("Add your API key in the left-hand panel first.")
        else:
            with st.spinner("Searching the regulations, answering, then checking the answer…"):
                ans = ai.ask(page, q.strip(), facts, LIB)
            S["chat"].append({"q": q.strip(), "a": ans, "page": page})
            S["chat_i"] = len(S["chat"]) - 1
    if S["chat"]:
        i = min(S["chat_i"], len(S["chat"]) - 1)
        item = S["chat"][i]
        c1, c2, c3 = st.columns([1, 3, 1])
        c1.button("◀  Older", key=f"older_{page}", disabled=i == 0, on_click=lambda: S.update(chat_i=i - 1), width="stretch")
        c2.markdown(f"<div style='text-align:center;padding-top:8px' class='muted'>Question {i + 1} of {len(S['chat'])}</div>",
                    unsafe_allow_html=True)
        c3.button("Newer  ▶", key=f"newer_{page}", disabled=i == len(S["chat"]) - 1, on_click=lambda: S.update(chat_i=i + 1),
                  width="stretch")
        st.markdown(f"**Q: {item['q']}**")
        show_ai(item["a"], "Answer" + ("" if item["a"].get("verified") else " · not fully confirmed by the checker"), f"chat_{page}")
        if item["a"].get("trail"):
            with st.expander("How this answer was checked"):
                for t in item["a"]["trail"]:
                    st.write("• " + t)


def nav_buttons():
    st.write("")
    st.divider()
    order = PAGES
    if S["page"] not in order:
        return
    i = order.index(S["page"])
    c1, _, c3 = st.columns([1, 2, 1])
    if i > 0:
        c1.button("←  Back", on_click=go, args=(order[i - 1],), width="stretch", key="back")
    if i < len(order) - 1:
        c3.button(f"Next: {order[i + 1].split('  ')[-1]}  →", on_click=go, args=(order[i + 1],), type="primary",
                  width="stretch", key="next", disabled=not ready())


def need_banks():
    st.info("First choose the deal on the **Home** page.")
    st.button("Go to Home", on_click=go, args=("Home",), type="primary")


# --------------------------------------------------------------------------------------
# Loading banks
# --------------------------------------------------------------------------------------
def parse(files):
    """Accepts the Fed's bulk .txt file and/or individual FR Y-9C PDFs (read exactly from their form fields)."""
    S.setdefault("files", {}).update(files)
    parsed, skipped = [], []
    for name, data in S["files"].items():
        q, banks = E.parse_y9c_pdf(data) if name.lower().endswith(".pdf") else E.parse_y9c_bulk(data)
        (parsed.append((q, banks)) if banks else skipped.append(name))
    if parsed:
        S["banks"] = E.merge_quarters(parsed)
    S["skipped"] = skipped


def add_test_files(names):
    files = {}
    for n in names:
        if n not in S.get("files", {}):
            with open(os.path.join(TEST, n), "rb") as fh:
                files[n] = fh.read()
    if files:
        parse(files)


def ihc_rssd(pk):
    return STRUCTURE[pk]["ihcs"][0]["rssd"] if STRUCTURE[pk]["ihcs"] else None


def load_example(kind):
    S["files"] = {}
    S["parent"], S["holdings"] = None, None
    S["saved"].pop("parent_src", None)
    add_test_files([SMALL_BANKS])
    by_name = {v["latest"]["name"]: k for k, v in S["banks"].items()}
    keep("tgt", by_name.get("ASSOCIATED BANC-CORP"))
    if kind == "canadian":
        keep("canadian", True)
        keep("parent_key", "RBC")
        add_test_files(US_ARM_FILES["RBC"])
        keep("acq", ihc_rssd("RBC"))
    else:
        keep("canadian", False)
        keep("tgt", by_name.get("WAFD, INC."))
        keep("acq", by_name.get("ASSOCIATED BANC-CORP"))


def buyer_label():
    return f"{saved('parent_key')}'s U.S. arm" if saved("canadian") else "Buyer"


# --------------------------------------------------------------------------------------
# The deal (everything later pages need)
# --------------------------------------------------------------------------------------
def deal():
    B = S["banks"]
    A, T = B[saved("acq")], B[saved("tgt")]
    a, t = A["latest"], T["latest"]
    c = E.combine(a, t)
    gsib = saved("gsib", False)
    cat_a, cat_t = E.category(A["avg_assets"], gsib), E.category(T["avg_assets"])
    pf_avg = A["avg_assets"] + T["avg_assets"]
    cat_c = E.category(pf_avg, gsib)
    quarters = sorted({b["quarter"] for X in (A, T) for b in X["history"] if b["quarter"]})
    return dict(A=A, T=T, a=a, t=t, c=c, cat_a=cat_a, cat_t=cat_t, cat_c=cat_c, pf_avg=pf_avg,
                liab=E.liabilities_test(c["total_liabilities"]), quarters=[nice_date(q) for q in quarters],
                min_q=min(A["quarters"], T["quarters"]))


def parent_values():
    P = S.get("parent")
    return {k: P["fig"][k]["value"] for k in P["fig"]} if P else {}


def canadian(d):
    """Everything the Canadian-buyer analysis needs. Pure calls to engine.py - no AI here."""
    pk = saved("parent_key")
    G = STRUCTURE[pk]
    v = parent_values()
    fx = saved("fx", 1.40)
    branches = saved("branches_bn", 0.0) * 1e6
    fp_ = E.us_footprint(d["a"]["total_assets"], branches, d["t"]["total_assets"])
    us_rows, (ihc_before, ihc_after) = E.us_rules_for_foreign_buyer(fp_)
    parent_total_usd = (v["total_assets"] * 1e3 / fx + d["t"]["total_assets"]) if v.get("total_assets") else None
    us_banking_after = fp_["combined_after"]
    q = E.qfbo_test(parent_total_usd, us_banking_after, saved("us_rev_share"), saved("us_ni_share"),
                    {"Yes": True, "No": False}.get(saved("mostly_banking", "Yes")))
    regk_rows = [
        {"status": q["test1"], "check": "Reg K test 1 — more than half its worldwide business is banking",
         "detail": "Confirmed by the user." if q["test1"] == "met" else "Not confirmed.", "src": "12 CFR 211.23(a)(1)"},
        {"status": q["test2"], "check": "Reg K test 2 — more than half its banking is outside the U.S. (after the deal)",
         "detail": "; ".join(f"{l['measure']}: {l['detail']}" for l in q["legs"]), "src": "12 CFR 211.23(a)(2)"}]
    holdings = list(S.get("holdings") or [])
    fed_recorded = [dict(h, source="Fed records") for h in G["foreign_nonbank"] + G["foreign_with_us_business"]]
    flags = E.holdings_flags({h["name"]: h for h in fed_recorded + holdings}.values())
    hold_rows = []
    n_flag = sum(1 for f in flags if f["flag"] != "met")
    hold_rows.append({"status": "review" if n_flag else ("met" if flags else "unknown"),
                      "check": "Holdings outside the U.S. that depend on Reg K status",
                      "detail": f"{n_flag} of {len(flags)} reviewed holdings need checking." if flags else
                      "No holdings list loaded — add the parent's list of subsidiaries in Step 4.", "src": "12 CFR 211.23(f)"})
    hold_rows.append({"status": "review" if G["outside_ihc"] else "met", "check": "U.S. companies held outside the IHC",
                      "detail": (", ".join(nice_name(x["name"]) for x in G["outside_ihc"][:3]) + " — confirm an exception applies or move under the IHC.")
                      if G["outside_ihc"] else "None — every U.S. company is under the IHC.", "src": "12 CFR 252.153"})
    osfi_rows = []
    if v.get("cet1"):
        to_cad = lambda usd_k: usd_k / 1e3 * fx
        parent = {k: v.get(k) for k in ("cet1", "tier1", "total_capital", "rwa", "leverage_exposure")}
        for r in E.parent_tests(parent, to_cad(d["t"]["rwa"]), to_cad(d["t"]["total_assets"]), saved("dsib", True), saved("dsb", 3.0)):
            osfi_rows.append({"status": r["status"], "check": f"Parent's {r['test']} (OSFI)",
                              "detail": f"{pct(r['before'])} → {pct(r['after'])} vs target {r['target']:.2f}%", "src": r["cite"]})
    else:
        osfi_rows.append({"status": "unknown", "check": "Parent's capital ratios (OSFI)", "detail": "Load the parent's Pillar 3 report on Home.",
                          "src": "OSFI CAR (2027) Ch.1"})
    return dict(pk=pk, G=G, v=v, fp=fp_, us_rows=us_rows, ihc_before=ihc_before, ihc_after=ihc_after, q=q, regk_rows=regk_rows,
                flags=flags, hold_rows=hold_rows, osfi_rows=osfi_rows, parent_total_usd=parent_total_usd,
                rows=us_rows + regk_rows + hold_rows + osfi_rows)


def summary_rows(d):
    rows = []
    before, after = d["cat_a"][0], d["cat_c"][0]
    rows.append({"status": "newly_applies" if E.rank(after) > E.rank(before) else "met", "check": "Size category of the combined U.S. bank",
                 "detail": f"{E.cat_label(before)} → {E.cat_label(after)} (combined {E.bn(d['pf_avg'])})", "src": E.CATEGORY_CITE})
    for r in E.newly_applying(before, after):
        rows.append({"status": "newly_applies", "check": r["Requirement"], "detail": f"{r['Before']} → {r['After']}", "src": r["Source"]})
    for r in E.capital_tests(d["c"]) + E.leverage_tests(d["c"], after):
        rows.append({"status": r["status"], "check": f"{r['test']} (combined)",
                     "detail": f"{pct(r['value'])} vs {r['with_buffer']:.1f}% needed" if r["value"] is not None else
                     "Not reported by both banks — can't combine", "src": r["cite"]})
    rows.append({"status": d["liab"]["status"], "check": "10% liabilities concentration limit", "detail": d["liab"]["detail"],
                 "src": "12 U.S.C. 1852(b); 12 CFR 251"})
    if saved("canadian"):
        cx = canadian(d)
        rows += [r for r in cx["us_rows"] if not r["check"].startswith("The IHC's own")] + cx["regk_rows"] + cx["hold_rows"] + cx["osfi_rows"]
    if d["min_q"] < 4:
        rows.append({"status": "unknown", "check": "Four-quarter average size",
                     "detail": f"Only {d['min_q']} quarter(s) loaded for at least one bank; size uses what is available.", "src": "12 CFR 252.5"})
    return rows


def show_summary_table(rows):
    table([{"Status": pill(r["status"]), "Check": r["check"], "Detail": r["detail"], "Source": r["src"]} for r in rows],
          ["Status", "Check", "Detail", "Source"], src_col="Source")


def counts(rows):
    n = {k: 0 for k in E.STATUS}
    for r in rows:
        n[r["status"]] += 1
    cols = st.columns(4)
    groups = [("met", ("met",), "Met"), ("not_met", ("not_met",), "Not met"),
              ("newly_applies", ("newly_applies", "warning"), "Newly applies / in buffer"), ("unknown", ("unknown", "review"), "Can't tell / review")]
    for col, (k, members, name) in zip(cols, groups):
        icon, _, colour = E.STATUS[k]
        col.markdown(f'<div class="stat"><div class="n" style="color:{colour}">{sum(n[m] for m in members)}</div>'
                     f'<div class="l">{icon} {name}</div></div>', unsafe_allow_html=True)
    st.write("")


# ======================================================================================
# PAGES
# ======================================================================================
def read_pillar3(src_bytes, src_name):
    if saved("parent_src") == fp(src_name, len(src_bytes)):
        return
    from pypdf import PdfReader
    text, pages = ai.find_capital_pages(PdfReader(io.BytesIO(src_bytes)))
    out = run_ai(fp("extract", src_name, len(src_bytes)), lambda: {"fig": ai.extract_parent_figures("Home", text), "pages": pages},
                 "The AI is reading the report's key-metrics table…")
    if out:
        keep("parent_src", fp(src_name, len(src_bytes)))
        S["parent"] = out


def page_home():
    st.markdown("# Bank merger regulatory check")
    st.markdown('<p class="lede">Choose the bank being bought and the buyer. The app shows which U.S. capital rules each one meets, '
                'what the merger changes — and, if the buyer is a Canadian bank, what U.S. regulation it takes on and what that means '
                'for the holdings it keeps at home.</p>', unsafe_allow_html=True)
    st.write("")
    left, right = st.columns([1.25, 1], gap="large")
    with left:
        st.markdown("### What you'll see")
        for n, title, body in [(1, "At a glance", "Every requirement — met, not met, newly applies."),
                               (2, "Capital ratios", "Each bank, and the two combined, against what's required."),
                               (3, "Leverage & liabilities", "The leverage ratio and the 10% liabilities cap."),
                               (4, "Canadian buyer", "The U.S. rules it takes on, Reg K, and the holdings to review."),
                               (5, "Summary & next steps", "A clean table, next steps and a draft memo.")]:
            st.markdown(f'<div class="step-preview"><div class="num">{n}</div><div><b style="font-size:17.5px">{title}</b><br>'
                        f'<span class="muted" style="font-size:15.5px">{body}</span></div></div>', unsafe_allow_html=True)
        st.write("")
        st.markdown("### How it works")
        st.markdown('<div class="howto"><div class="ic">🧮</div><div><b>The numbers are calculated</b>'
                    '<span>Every figure and every met / not met result follows fixed rules taken from the regulations — '
                    'the same answer every time.</span></div></div>'
                    '<div class="howto"><div class="ic">✨</div><div><b>The explanations are written by AI</b>'
                    '<span>Using only the regulations loaded into the app, with the sources shown so you can check them.</span></div></div>',
                    unsafe_allow_html=True)
    with right:
        st.markdown("### Choose the deal")
        with st.popover("Upload the banks' FR Y-9C reports", width="stretch"):
            st.markdown("From the Fed's [National Information Center](https://www.ffiec.gov/npw): search the bank → "
                        "**Financial Reports** → **FR Y-9C** (PDF), or the bulk **Financial Data Download** (.txt). Up to four quarters.")
            files = st.file_uploader("FR Y-9C files", type=["txt", "pdf"], accept_multiple_files=True, label_visibility="collapsed")
            if files:
                parse({f.name: f.getvalue() for f in files})
                if S.get("skipped"):
                    st.warning("Not an FR Y-9C report, so skipped: " + ", ".join(S["skipped"]))
        c0, c1, c2 = st.columns([1.05, 1.15, 0.8], gap="small")
        c0.markdown('<div class="muted" style="padding-top:6px">Or try an example:</div>', unsafe_allow_html=True)
        if c1.button("Canadian buyer", type="tertiary", help="RBC buying Associated Banc-Corp."):
            load_example("canadian")
        if c2.button("U.S. buyer", type="tertiary", help="Associated Banc-Corp buying WaFd."):
            load_example("us")
        if not S.get("banks"):
            st.markdown('<span class="muted">Upload the two banks\' reports, or try an example.</span>', unsafe_allow_html=True)
            return
        B = S["banks"]
        arm_ids = {ihc_rssd(k) for k in CANADIAN}
        targets = sorted([k for k in B if k not in arm_ids], key=lambda k: -B[k]["latest"]["total_assets"])
        lab = lambda k: f"{nice_name(B[k]['latest']['name'])}  ·  {E.bn(B[k]['latest']['total_assets'])}"
        if not targets:
            st.warning("Upload at least one U.S. bank to be bought.")
            return
        tgt = st.selectbox("Bank being bought (U.S.)", targets, format_func=lab,
                           index=targets.index(saved("tgt")) if saved("tgt") in targets else 0)
        keep("tgt", tgt)
        is_can = st.checkbox("The buyer is a Canadian bank", value=saved("canadian", False),
                             help="TD, RBC, BMO or CIBC buying a U.S. bank to grow its U.S. presence.")
        keep("canadian", is_can)
        if is_can:
            with st.container(border=True):
                pk = st.selectbox("Which Canadian bank?", CANADIAN, format_func=lambda k: f"{nice_name(STRUCTURE[k]['parent'])} ({k})",
                                  index=CANADIAN.index(saved("parent_key")) if saved("parent_key") in CANADIAN else 0)
                if pk != saved("parent_key"):
                    S["parent"], S["holdings"] = None, None
                    S["saved"].pop("parent_src", None)
                keep("parent_key", pk)
                add_test_files(US_ARM_FILES.get(pk, []))
                keep("acq", ihc_rssd(pk))
                st.markdown(f'Its U.S. arm: <b>{nice_name(STRUCTURE[pk]["ihcs"][0]["name"])}</b> '
                            f'({E.bn(S["banks"][ihc_rssd(pk)]["latest"]["total_assets"])}, FR Y-9C).', unsafe_allow_html=True)
                st.markdown("**Its Pillar 3 report** (for its Canadian capital ratios)")
                c1, c2 = st.columns([1, 1.4])
                if pk in PILLAR3 and c1.button("Use example", width="stretch", help=PILLAR3[pk][1]):
                    with open(os.path.join(TEST, PILLAR3[pk][0]), "rb") as fh:
                        read_pillar3(fh.read(), PILLAR3[pk][0])
                pdf = c2.file_uploader("Pillar 3 PDF", type=["pdf"], label_visibility="collapsed", key="p3_upload")
                if pdf is not None:
                    read_pillar3(pdf.getvalue(), pdf.name)
                v = parent_values()
                if v:
                    chk = ai.verify_parent(S["parent"]["fig"])
                    ok = all(c["ok"] for c in chk if c["ok"] is not None) and any(c["ok"] for c in chk)
                    verdict = "✓ double-checked against the bank's own ratios" if ok else "? some figures missing — check them in Step 4"
                    st.markdown(f'<span class="ai-tag">✨ Read by AI</span>&nbsp; {v.get("period") or "Latest quarter"}: CET1 '
                                f'C${(v.get("cet1") or 0):,.0f}m, total assets C${(v.get("total_assets") or 0):,.0f}m · '
                                f'<span class="code-tag">{verdict}</span>', unsafe_allow_html=True)
        else:
            buyers = [k for k in targets if k != tgt] or targets
            acq = st.selectbox("Buyer (U.S.)", buyers, format_func=lab, index=buyers.index(saved("acq")) if saved("acq") in buyers else 0)
            keep("acq", acq)
            keep("gsib", st.checkbox("The buyer is a U.S. GSIB", value=saved("gsib", False)))
        if ready():
            st.write("")
            st.button("Get started  →", type="primary", on_click=go, args=(PAGES[1],), width="stretch")
    if ready():
        d = deal()
        st.write("")
        with st.expander("AI's first read of the two banks' figures"):
            for b in (d["a"], d["t"]):
                figs = "\n".join(f"- {E.FIELDS[k][1]}: {pct(b[k]) if k in E.RATIO_FIELDS else E.bn(b[k], 2)}" for k in E.FIELDS)
                checks = "\n".join(f"- {c['check']}: {c['result']}" for c in E.cross_checks(b))
                res = run_ai(fp("review", b["rssd"], b["quarter"]), lambda: {"text": ai.review_bank("Home", b["name"], figs, checks)},
                             f"Reading {nice_name(b['name'])}…")
                st.markdown(f"**{nice_name(b['name'])}**")
                show_ai(res, "figures only — no law", f"review_{b['rssd']}")


def page_glance():
    st.markdown("## 1 · At a glance")
    if not ready():
        return need_banks()
    d = deal()
    st.markdown(f'<p class="lede"><b>{nice_name(d["a"]["name"])}</b> buying <b>{nice_name(d["t"]["name"])}</b> — '
                f'combined average assets {E.bn(d["pf_avg"])}.</p>', unsafe_allow_html=True)
    rows = summary_rows(d)
    counts(rows)
    show_summary_table(rows)
    st.markdown('<span class="muted">Combined bank = the two banks added together (no purchase premium or goodwill). '
                'Size categories use asset size only.</span>', unsafe_allow_html=True)
    facts = "\n".join(f"- {E.STATUS[r['status']][1]} | {r['check']} | {r['detail']}" for r in rows)
    st.write("")
    res = run_ai(fp("glance", facts), lambda: ai.explain(
        "1 At a glance", "In 3-4 sentences, give the deal lawyer the headline: what the merger changes in regulatory terms and the "
        "single most important item to watch. Cite passages.", facts,
        "Category IV $100 billion average total consolidated assets capital plan stress capital buffer minimum capital ratios", LIB))
    show_ai(res, "Headline", "glance")
    chat_box("1 At a glance", facts)


def ratio_cell(value, status):
    icon, _, colour = E.STATUS[status]
    return f'<span style="color:{colour};font-weight:600">{icon}&nbsp; {pct(value)}</span>'


def page_capital():
    st.markdown("## 2 · Capital ratios")
    if not ready():
        return need_banks()
    d = deal()
    st.markdown('<p class="lede">Every bank must keep three ratios of capital to risk-weighted assets above a minimum, '
                'plus a 2.5% buffer. Below the buffer, dividends and bonuses are restricted; below the minimum, the bank is '
                'under-capitalised.</p>', unsafe_allow_html=True)
    rows = []
    for k, name, m in E.CAPITAL_RATIOS:
        r = {"Ratio": f"<b>{name}</b>", "Needed": f"{m:.1f}% minimum<br><span class='muted'>{m + E.BUFFER:.1f}% with buffer</span>"}
        for col, bank in (("Buyer", d["a"]), ("Bank bought", d["t"]), ("Combined", d["c"])):
            r[col] = ratio_cell(bank[k], E.judge(bank[k], m, m + E.BUFFER))
        rows.append(r)
    table(rows, ["Ratio", "Needed", "Buyer", "Bank bought", "Combined"])
    st.markdown(f'<span class="muted">Buyer: {nice_name(d["a"]["name"])}{" (" + saved("parent_key") + "\'s U.S. arm)" if saved("canadian") else ""} · Bank bought: {nice_name(d["t"]["name"])} · '
                f'Source: 12 CFR 217.10(a) (minimums), 217.11(a) (buffer). For Category IV and above, the buffer is the bank\'s '
                f'stress capital buffer, which is at least 2.5% (12 CFR 225.8).</span>', unsafe_allow_html=True)

    with st.expander("Where these numbers come from"):
        st.markdown('<span class="code-tag">✓ Calculated</span>&nbsp; Read straight from the Fed filings and double-checked.', unsafe_allow_html=True)
        frows = []
        for k in ("cet1", "tier1", "total_capital", "rwa"):
            codes, lbl, loc = E.FIELDS[k]
            frows.append({"Figure": lbl, "Buyer": E.bn(d["a"][k]), "Bank bought": E.bn(d["t"][k]), "Combined": E.bn(d["c"][k]),
                          "FR Y-9C line": f"{' + '.join(codes)} · {loc}"})
        table(frows, ["Figure", "Buyer", "Bank bought", "Combined", "FR Y-9C line"], src_col="FR Y-9C line")
        for b in (d["a"], d["t"]):
            st.markdown(f"**{nice_name(b['name'])}** — " + " · ".join(
                f"{'✓' if c['ok'] else '?' if c['ok'] is None else '✕'} {c['check']}" for c in E.cross_checks(b)))

    facts = "\n".join(f"{name}: buyer {pct(d['a'][k])}, bank bought {pct(d['t'][k])}, combined {pct(d['c'][k])}; needed {m}% min / {m + E.BUFFER}% with buffer"
                      for k, name, m in E.CAPITAL_RATIOS) + f"\nCombined category: {E.cat_label(d['cat_c'][0])}"
    res = run_ai(fp("capital", facts), lambda: ai.explain(
        "2 Capital ratios", "In under 150 words, explain how the buyer, the bank bought, and the combined bank meet the capital "
        "ratio requirements, and which ratio has the least room. Cite passages.", facts,
        "minimum capital requirements common equity tier 1 capital ratio tier 1 total capital capital conservation buffer", LIB))
    show_ai(res, "What this means", "capital")
    chat_box("2 Capital ratios", facts)


def page_leverage():
    st.markdown("## 3 · Leverage & liabilities")
    if not ready():
        return need_banks()
    d = deal()
    cat = d["cat_c"][0]
    st.markdown("### Leverage")
    st.markdown('<p class="lede">The leverage ratio compares Tier 1 capital with total assets, ignoring risk. '
                'Every bank needs 4%. From Category III, a second "supplementary" ratio of 3% also applies.</p>', unsafe_allow_html=True)
    rows = []
    for r in E.leverage_tests(d["c"], cat):
        row = {"Ratio": f"<b>{r['test']}</b>", "Needed": f"{r['minimum']:.1f}%"}
        for col, bank in (("Buyer", d["a"]), ("Bank bought", d["t"]), ("Combined", d["c"])):
            v = bank.get(r["key"])
            row[col] = ratio_cell(v, E.judge(v, r["minimum"], r["minimum"])) if v is not None else "<span class='muted'>not reported</span>"
        rows.append(row)
    table(rows, ["Ratio", "Needed", "Buyer", "Bank bought", "Combined"])
    st.markdown('<span class="muted">Source: 12 CFR 217.10(a)(1)(iv) and (v).</span>', unsafe_allow_html=True)

    st.markdown("### Liabilities")
    st.markdown('<p class="lede">No merger may leave the combined firm with more than 10% of the liabilities of the whole U.S. '
                'financial sector.</p>', unsafe_allow_html=True)
    L = d["liab"]
    share_of_limit = min(1.0, d["c"]["total_liabilities"] / L["limit"])
    c1, c2 = st.columns([3, 1])
    with c1:
        st.progress(share_of_limit, text=(f"Combined liabilities {E.bn(d['c']['total_liabilities'])} — {share_of_limit * 100:.1f}% "
                                          f"of the way to the {E.bn(L['limit'])} limit").replace("$", "\\$"))
    c2.markdown(pill(L["status"]), unsafe_allow_html=True)
    st.markdown(f'<span class="muted">Source: 12 U.S.C. 1852(b); 12 CFR 251. Sector total {E.bn(E.FIN_SECTOR_LIABILITIES)} — '
                f'{E.FIN_SECTOR_SOURCE}</span>', unsafe_allow_html=True)

    facts = (f"Leverage ratio: buyer {pct(d['a']['leverage_ratio'])}, bank bought {pct(d['t']['leverage_ratio'])}, combined "
             f"{pct(d['c']['leverage_ratio'])} vs 4% minimum. Combined category {E.cat_label(cat)} (SLR applies from Category III). "
             f"Liabilities: {L['detail']}")
    st.write("")
    res = run_ai(fp("leverage", facts), lambda: ai.explain(
        "3 Leverage & liabilities", "In under 150 words explain the leverage position and the liabilities concentration limit for "
        "this deal. Cite passages.", facts,
        "leverage ratio 4 percent supplementary leverage ratio concentration limit consolidated liabilities 10 percent financial companies", LIB))
    show_ai(res, "What this means", "leverage")
    chat_box("3 Leverage & liabilities", facts)


def page_canadian():
    st.markdown("## 4 · Canadian buyer")
    if not ready():
        return need_banks()
    if not saved("canadian"):
        st.markdown('<p class="lede">This step is for a Canadian bank buying a U.S. bank. Tick <b>"The buyer is a Canadian bank"</b> '
                    'on the Home page to use it.</p>', unsafe_allow_html=True)
        st.button("Go to Home", on_click=go, args=("Home",))
        return
    d = deal()
    cx = canadian(d)
    G, v, fp_ = cx["G"], cx["v"], cx["fp"]
    parent_name = nice_name(G["parent"])
    ihc_name = nice_name(G["ihcs"][0]["name"]) if G["ihcs"] else "its U.S. holding company"
    st.markdown(f'<p class="lede">Once a foreign bank\'s U.S. presence is big enough, the U.S. regulates it much like a U.S. bank — '
                f'and that reaches back to the businesses it owns at home. Here: <b>{parent_name}</b> buying '
                f'<b>{nice_name(d["t"]["name"])}</b>.</p>', unsafe_allow_html=True)

    # ---- A. U.S. footprint ----
    st.markdown("### A. Its U.S. presence, before and after")
    st.markdown(f'<span class="muted">From {ihc_name}\'s FR Y-9C and the Fed\'s ownership records.</span>', unsafe_allow_html=True)
    c1, c2, c3 = st.columns(3)
    for col, lbl, before, after in ((c1, "U.S. subsidiaries (non-branch)", fp_["nonbranch_before"], fp_["nonbranch_after"]),
                                    (c2, "All U.S. operations incl. branches", fp_["combined_before"], fp_["combined_after"])):
        col.markdown(f'<div class="stat"><div class="l">{lbl}</div><div class="n" style="font-size:22px">{E.bn(before)} → {E.bn(after)}</div></div>',
                     unsafe_allow_html=True)
    c3.markdown(f'<div class="stat"><div class="l">Today in the U.S.</div><div class="n" style="font-size:22px">{G["under_ihc"]:,} companies</div>'
                f'<div class="l">under the IHC · {len(G["branches"])} branches/offices</div></div>', unsafe_allow_html=True)
    st.write("")
    keep("branches_bn", st.number_input("Assets in its U.S. branches and agencies (US$bn, optional)", 0.0, value=float(saved("branches_bn", 0.0)),
                                        step=1.0, help="From the branches' FFIEC 002 reports (available from the Fed on request). "
                                        "Branches count toward 'all U.S. operations' but not toward the $50bn IHC trigger."))

    # ---- B. U.S. rules ----
    st.markdown("### B. The U.S. rules it is now subject to")
    st.markdown(f'<span class="muted">{E.SUBPART_O_NOTE}</span>', unsafe_allow_html=True)
    show_summary_table(cx["us_rows"])
    new_ihc = E.newly_applying(cx["ihc_before"], cx["ihc_after"])
    if new_ihc:
        st.markdown(f"**What {ihc_name} newly has to do** ({E.cat_label(cx['ihc_before'])} → {E.cat_label(cx['ihc_after'])}):")
        table([{"Requirement": r["Requirement"], "Before": r["Before"], "After": r["After"], "Source": r["Source"]} for r in new_ihc],
              ["Requirement", "Before", "After", "Source"], src_col="Source")

    # ---- C. Reg K ----
    st.markdown("### C. Reg K — can it keep its businesses at home?")
    st.markdown('<p class="lede">A foreign bank may run <b>any</b> business outside the U.S. — and own foreign companies that also do some '
                'U.S. business — only while it is a <b>qualifying foreign banking organization</b>. It must stay mostly a bank, and '
                'mostly outside the U.S. A big U.S. acquisition pushes on the second test.</p>', unsafe_allow_html=True)
    c1, c2, c3 = st.columns(3)
    keep("mostly_banking", c1.radio("Is more than half its worldwide business banking?", ["Yes", "No", "Not sure"],
                                    index=["Yes", "No", "Not sure"].index(saved("mostly_banking", "Yes")), horizontal=True,
                                    help="Test 1. True for the big Canadian banks; confirm from its annual FR Y-7 report to the Fed."))
    keep("us_rev_share", c2.number_input("Share of its banking revenue earned in the U.S. after the deal (%)", 0.0, 100.0,
                                         value=saved("us_rev_share"), step=1.0, help="From its annual report's U.S. segment, plus the target."))
    keep("us_ni_share", c3.number_input("Share of its banking net income earned in the U.S. after the deal (%)", 0.0, 100.0,
                                        value=saved("us_ni_share"), step=1.0, help="From its annual report's U.S. segment, plus the target."))
    cx = canadian(d)
    q = cx["q"]
    table([{"Measure": l["measure"], "Share in the U.S.": "n/a" if l["us_share"] is None else f"{l['us_share']:.1f}%",
            "Result": pill({True: "met", False: "not_met", None: "unknown"}[l["pass"]]), "Detail": l["detail"]} for l in q["legs"]],
          ["Measure", "Share in the U.S.", "Result", "Detail"])
    st.markdown(f"**Overall: {pill(q['overall'])}** &nbsp; Test 1 {pill(q['test1'])} &nbsp; Test 2 {pill(q['test2'])}"
                + (f" &nbsp; · &nbsp; Room for more U.S. banking assets before the assets measure fails: <b>{E.bn(q['headroom'])}</b>"
                   if q["headroom"] else ""), unsafe_allow_html=True)
    st.markdown('<span class="muted">Each test needs 2 of 3 measures. Assets come from the Pillar 3 report and FR Y-9Cs (U.S. assets include '
                'branches entered above); Puerto Rico, Guam and other territories count as the U.S. The parent\'s total assets are used as '
                '"outside the U.S." after removing U.S. assets — a simplification. Source: 12 CFR 211.23(a)–(b).</span>', unsafe_allow_html=True)
    with st.expander("What happens if it stops qualifying"):
        for title, body, cite in E.REG_K_CONSEQUENCES:
            st.markdown(f"**{title}.** {body} <span class='muted'>({cite})</span>", unsafe_allow_html=True)

    # ---- D. Holdings ----
    st.markdown("### D. Holdings to review")
    st.markdown("**Outside the U.S.** — businesses that only stay permissible while it keeps Reg K status, or that also reach into the U.S.")
    with st.container(border=True):
        st.markdown("Add the parent's list of subsidiaries (the **principal subsidiaries** page of its annual report, as a PDF, or pasted "
                    "text). AI sorts each company by type of business; the app then applies the Reg K rules to each one.", unsafe_allow_html=True)
        c1, c2 = st.columns(2)
        hpdf = c1.file_uploader("Annual report or subsidiaries list (PDF)", type=["pdf"], key="holdings_pdf")
        htxt = c2.text_area("…or paste the list", height=110, key="holdings_txt", placeholder="Company name — country — business")
        if st.button("Review these holdings", type="primary"):
            if not (hpdf or htxt.strip()):
                st.warning("Upload a PDF or paste a list first.")
                st.stop()
            if hpdf is not None:
                from pypdf import PdfReader
                text, pages = ai.find_subsidiary_pages(PdfReader(io.BytesIO(hpdf.getvalue())))
            else:
                text = htxt
            res = run_ai(fp("holdings", cx["pk"], text[:2000], len(text)),
                         lambda: {"h": ai.classify_holdings("4 Canadian buyer", text, parent_name)}, "The AI is sorting the holdings…")
            if res:
                S["holdings"] = res["h"]
                st.rerun()
    cx = canadian(d)
    if cx["flags"]:
        table([{"Flag": pill(f["flag"]), "Company": nice_name(f["name"]), "Country": f["country"], "Business": f["activity"],
                "Why": f["why"], "From": f.get("source", "Your list")} for f in cx["flags"]],
              ["Flag", "Company", "Country", "Business", "Why", "From"], src_col="From")
    st.markdown(f'<span class="muted">The Fed\'s records list only the foreign companies {parent_name} reports to the Fed '
                f'({G["foreign_total"]} here), so add its own list for a full review.</span>', unsafe_allow_html=True)
    if G.get("territory_holdings"):
        st.markdown(f'<span class="muted">{len(G["territory_holdings"])} companies in U.S. territories (e.g. Puerto Rico) are treated as '
                    f'<b>U.S.</b> under Reg K.</span>', unsafe_allow_html=True)

    st.markdown("**In the U.S.** — everything must sit under the IHC, except branches and narrow exceptions.")
    rows = [{"Item": f"<b>{nice_name(d['t']['name'])}</b> (being bought)", "What to do": "Place under the IHC.", "Source": "12 CFR 252.153"}]
    for x in G["outside_ihc"]:
        rows.append({"Item": nice_name(x["name"]) + f" <span class='muted'>({x['type']}, {x['where']})</span>",
                     "What to do": "Held outside the IHC today — confirm it is a section 2(h)(2) company or DPC branch subsidiary, or move it.",
                     "Source": "12 CFR 252.153"})
    rows.append({"Item": f"{len(G['branches'])} U.S. branches, agencies and representative offices",
                 "What to do": "Stay outside the IHC (part of the parent itself). Branch vs subsidiary is a structuring choice.",
                 "Source": "12 CFR 252.153"})
    table(rows, ["Item", "What to do", "Source"], src_col="Source")
    with st.expander(f"Its U.S. branches and offices ({len(G['branches'])})"):
        table([{"Office": nice_name(b["name"]), "Type": b["type"], "Where": b["where"]} for b in G["branches"]], ["Office", "Type", "Where"])

    # ---- E. OSFI ----
    st.markdown("### E. Its Canadian capital (OSFI)")
    P = S.get("parent")
    if P:
        fig = P["fig"]
        st.markdown(f'<span class="ai-tag">✨ Read by AI</span>&nbsp; from pages {", ".join(map(str, P["pages"]))} of the Pillar 3 report. '
                    f'You can correct any figure in the table.', unsafe_allow_html=True)
        keys = ai.AMOUNT_FIELDS + ("cet1_ratio", "leverage_ratio")
        edit = pd.DataFrame([{"Figure": ai.PARENT_FIELDS[k].split(" (")[0], "C$ millions / %": fig[k]["value"],
                              "Where the AI found it": fig[k]["found_as"]} for k in keys])
        edited = st.data_editor(edit, hide_index=True, width="stretch", disabled=["Figure", "Where the AI found it"], key="parent_editor")
        for k, val in zip(keys, edited["C$ millions / %"]):
            fig[k]["value"] = None if pd.isna(val) else float(val)
        st.markdown('<span class="code-tag">✓ Double-check</span>&nbsp; ' + " · ".join(
            f"{'✓' if c['ok'] else '?' if c['ok'] is None else '✕'} {c['check']} ({c['result']})" for c in ai.verify_parent(fig)),
            unsafe_allow_html=True)
    c1, c2, c3 = st.columns(3)
    keep("fx", c1.number_input("C\\$ per US\\$1", 0.5, 3.0, float(saved("fx", 1.40)), 0.01, help="Enter today's exchange rate."))
    keep("dsib", c2.checkbox("It is a Canadian D-SIB", value=saved("dsib", True), help="TD, RBC, BMO and CIBC all are."))
    keep("dsb", c3.number_input("OSFI Domestic Stability Buffer (%)", 0.0, 4.0, float(saved("dsb", 3.0)), 0.25,
                                help="Set by OSFI; not in this app's library. Check OSFI's latest announcement."))
    cx = canadian(d)
    show_summary_table(cx["osfi_rows"])
    st.markdown('<span class="muted">Simplified: it pays cash at book value, so its capital is unchanged and the target\'s risk-weighted assets '
                'are added at their U.S. amount (OSFI\'s weights may differ).</span>', unsafe_allow_html=True)

    facts = "\n".join(f"- {E.STATUS[r['status']][1]} | {r['check']} | {r['detail']}" for r in cx["rows"])
    facts += f"\nParent: {G['parent']}; U.S. holding company: {ihc_name}."
    res = run_ai(fp("canadian", facts), lambda: ai.explain(
        "4 Canadian buyer", "In under 200 words, explain to the deal lawyer: which U.S. rules the Canadian bank is (newly) subject to after "
        "the deal; whether it still qualifies under Reg K to keep its non-U.S. businesses, and what happens if not; and which holdings "
        "need review. Use the statuses exactly as given. Cite passages.",
        facts, "qualifying foreign banking organization more than half banking outside United States nonbanking activities exemptions "
        "U.S. intermediate holding company combined U.S. assets $100 billion risk committee", LIB, k=8))
    show_ai(res, "What this means", "canadian")
    chat_box("4 Canadian buyer", facts)


NEXT_STEPS = {
    "Size category": "Plan for the new category's requirements to be in place by closing — supervisors expect it.",
    "Annual capital plan (CCAR)": "Start building the capital plan and stress-testing capability (12 CFR 225.8).",
    "Supervisory stress test": "Prepare for the Fed's supervisory stress test cycle.",
    "Supplementary leverage ratio (3%)": "Model the combined supplementary leverage ratio using both banks' leverage exposure.",
    "Liquidity coverage ratio / NSFR": "Check whether the combined bank's funding mix triggers the liquidity coverage ratio.",
    "Hold all U.S. subsidiaries": "Place the bank being bought under the U.S. intermediate holding company.",
    "U.S. companies held outside the IHC": "Confirm each U.S. company outside the IHC falls within an exception, or move it under the IHC.",
    "U.S. risk committee": "Set up the U.S. risk committee and appoint a U.S. chief risk officer (12 CFR 252.155).",
    "Liquidity stress tests": "Build U.S. liquidity stress testing and hold a U.S. liquidity buffer (12 CFR 252.157).",
    "Capital stress testing": "Prepare for capital stress testing of the U.S. operations (12 CFR 252.158).",
    "Certify the parent": "Prepare the certification that the parent meets Basel-consistent home capital standards (12 CFR 252.154).",
    "Reg K test": "Confirm the Reg K tests with revenue and net income from the annual report; consider asking the Fed for a specific determination (12 CFR 211.23(e)).",
    "Holdings outside the U.S.": "Review each flagged holding outside the U.S. against Reg K (12 CFR 211.23(f)).",
    "Parent's capital ratios": "Load the parent's Pillar 3 report to check its OSFI ratios.",
    "Four-quarter average size": "Upload the other three quarters of FR Y-9C data to confirm the size category.",
}


def page_summary():
    st.markdown("## 5 · Summary & next steps")
    if not ready():
        return need_banks()
    d = deal()
    rows = summary_rows(d)
    st.markdown(f'<p class="lede"><b>{nice_name(d["a"]["name"])}</b> buying <b>{nice_name(d["t"]["name"])}</b> · '
                f'data as of {", ".join(d["quarters"])}</p>', unsafe_allow_html=True)
    counts(rows)
    show_summary_table(rows)

    st.markdown("### Next steps")
    steps = []
    for r in rows:
        if r["status"] in ("met", "applies"):
            continue
        text = next((t for k, t in NEXT_STEPS.items() if r["check"].startswith(k)), None)
        if text is None and r["status"] in ("not_met", "warning"):
            text = f"{r['detail']}. Model a capital raise or a smaller deal before signing."
        if text:
            steps.append(f"**{r['check']}** — {text}")
    if not steps:
        steps = ["Nothing is flagged on the figures loaded. Confirm with all four quarters of data before relying on it."]
    for i, s_ in enumerate(dict.fromkeys(steps), 1):
        st.markdown(f"{i}. {s_}")

    facts = f"Deal: {d['a']['name']} buying {d['t']['name']}.\n" + "\n".join(
        f"- {E.STATUS[r['status']][1]} | {r['check']} | {r['detail']} | {r['src']}" for r in rows)
    st.markdown("### Draft memo")

    def memo():
        draft = ai.explain("5 Summary", "Draft a short internal memo (under 250 words) to a banking M&A partner: headline, items not met "
                           "or newly applying, items that could not be checked and what data is needed, next steps. Use the statuses "
                           "exactly as given; cite passages for legal points.", facts,
                           "Category IV capital plan stress capital buffer minimum capital ratios leverage ratio concentration limit "
                           "intermediate holding company", LIB, k=8)
        draft["verdict"] = ai.reflect("5 Summary", draft["text"], draft["passages"], "Is the memo supported by its sources?")
        return draft

    res = run_ai(fp("memo", facts), memo, "Drafting the memo, then a second AI call is checking it…")
    show_ai(res, "Draft memo", "memo")
    if res:
        verdict = res.get("verdict", {})
        if verdict.get("supported"):
            st.success("Checked by a second AI call: every legal statement is supported by its cited source.")
        elif verdict:
            flagged = [x for x in verdict.get("unsupported", []) if x][:3]
            st.warning("A second AI call could not confirm every statement against its sources"
                       + (": " + "; ".join(flagged) if flagged else ".") + " Treat the memo as a draft to check.")
    c1, c2 = st.columns(2)
    c1.download_button("Download as Word", make_docx(d, rows, res, steps), file_name="merger_regulatory_summary.docx",
                       type="primary", width="stretch")
    c2.download_button("Download the AI audit trail", json.dumps(S.get("audit", []), indent=2), file_name="ai_audit_trail.json",
                       width="stretch")
    chat_box("5 Summary", facts)


def make_docx(d, rows, memo, steps):
    from docx import Document
    doc = Document()
    doc.add_heading("Bank merger regulatory check", 0)
    doc.add_paragraph(f"{nice_name(d['a']['name'])} buying {nice_name(d['t']['name'])}. Data as of {', '.join(d['quarters'])}.")
    doc.add_heading("Results", 1)
    tbl = doc.add_table(rows=1, cols=4)
    tbl.style = "Light Grid Accent 1"
    for i, h in enumerate(["Status", "Check", "Detail", "Source"]):
        tbl.rows[0].cells[i].text = h
    for r in rows:
        cells = tbl.add_row().cells
        for i, val in enumerate([E.STATUS[r["status"]][1], r["check"], r["detail"], r["src"]]):
            cells[i].text = val
    doc.add_heading("Next steps", 1)
    for s_ in dict.fromkeys(steps):
        doc.add_paragraph(s_.replace("**", ""), style="List Number")
    if memo:
        doc.add_heading("Draft memo (AI-generated — verify before use)", 1)
        for para in memo["text"].split("\n"):
            if para.strip():
                doc.add_paragraph(para.strip().lstrip("#").strip())
    doc.add_paragraph("Screening tool, not legal advice. The combined bank is a simple sum of the two banks; verify against the filings and rules cited.")
    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()


def page_sizes():
    st.markdown("## Size examples")
    st.markdown('<p class="lede">Which category a merger lands in depends mainly on the combined bank\'s average total assets. '
                'These made-up pairs show each step of the ladder.</p>', unsafe_allow_html=True)
    examples = [
        ("Below $100bn — no category", "Harbor Bancorp ($45bn) + Prairie Financial ($30bn)", 75e6, False),
        ("Category IV", "Lakeshore Bancorp ($85bn) + Summit Bank ($40bn)", 125e6, False),
        ("Category III", "Pacific Crest Financial ($190bn) + Gulf Coast Bancshares ($85bn)", 275e6, False),
        ("Category II", "Atlantic National ($480bn) + Continental Trust ($260bn)", 740e6, False),
        ("Category I (GSIB)", "A U.S. GSIB ($2.5tn) + Riverbend Bank ($30bn)", 2530e6, True),
    ]
    rows = []
    for label, pair, total, gsib in examples:
        cat = E.category(total, gsib)[0]
        rows.append({"Lands in": f"<b>{label}</b>", "Example deal (fictitious)": pair, "Combined": E.bn(total),
                     "What switches on": ", ".join(r[0].split(" (")[0] for r in E.REQUIREMENTS if r[1][cat] not in ("No", "Optional opt-out")) or "Base rules only"})
    table(rows, ["Lands in", "Example deal (fictitious)", "Combined", "What switches on"])
    st.markdown('<span class="muted">There is no "Category 5": below $100bn a bank simply has no category. Category I is only for '
                'U.S. GSIBs, which are identified by the Fed\'s GSIB method, not size alone. Source: 12 CFR 252.5.</span>',
                unsafe_allow_html=True)
    st.markdown("### Try your own")
    c1, c2, c3 = st.columns(3)
    x = c1.number_input("Buyer ($bn)", 0.0, value=120.0, step=5.0)
    y = c2.number_input("Bank bought ($bn)", 0.0, value=60.0, step=5.0)
    g = c3.checkbox("Buyer is a GSIB")
    cat, why = E.category((x + y) * 1e6, g)
    st.markdown(f"**{E.cat_label(cat)}** — {why}")


# ======================================================================================
{"Home": page_home, PAGES[1]: page_glance, PAGES[2]: page_capital, PAGES[3]: page_leverage, PAGES[4]: page_canadian,
 PAGES[5]: page_summary, EXTRA: page_sizes}[S["page"]]()
nav_buttons()
counter.markdown(f'<span class="muted">AI calls this session: <b>{S.get("ai_calls", 0)}</b></span>', unsafe_allow_html=True)
